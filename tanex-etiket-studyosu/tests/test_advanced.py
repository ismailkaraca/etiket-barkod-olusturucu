"""Regression checks for Tanex Label Studio production features."""
import io
import sys
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from core import inspect_template,compose_labels,export_docx,export_pdf,export_calibration_pdf,validate_records
from docx import Document

TEMPLATE=(ROOT/'templates/TW-2408_105x72mm.docx').read_bytes()
META=inspect_template(TEMPLATE,'TW-2408_105x72mm.docx')
HEADERS=['Kütüphane Adı','Adres','İlçe','İl','Telefon','Paket','Etiket Adedi']
RECORDS=[{'id':2,'qty':1,'values':['Örnek Kütüphane','Atatürk Mahallesi','Sincan','Ankara','05XX XXX XX XX','1/2','2.000']}]

def test_shipment_preview_word_pdf():
    opts={'mode':'shipping','shippingMap':{'library':'0','address':'1','district':'2','city':'3','phone':'4','package':'5','count':'6'},'shippingTitle':'RFID ETİKET GÖNDERİSİ','padding':2}
    out=compose_labels(META,{'fields':[],'headers':HEADERS,'records':RECORDS,'options':opts,'start':0})
    assert out['total']==1 and out['pageCount']==1
    parts=out['pages'][0][0]['parts']
    assert parts[0]['value']=='RFID ETİKET GÖNDERİSİ'
    assert any(part['value']=='Sincan / Ankara' for part in parts)
    document=Document(io.BytesIO(export_docx(TEMPLATE,META,out,opts)))
    assert document.tables and 'Kütüphane' in document.tables[0].cell(0,0).text
    assert export_pdf(META,out,opts).startswith(b'%PDF')
    assert export_calibration_pdf(META,opts).startswith(b'%PDF')

def test_rich_content_and_codes():
    options={'padding':2,'fieldOptions':{'0':{'showHeading':True,'label':'Kütüphane','boldHeading':True,'scale':1.3,'color':'#123456'}},
             'extraFields':[{'kind':'text','formula':'{İlçe} / {İl}'},{'kind':'qr','formula':'{Kütüphane Adı}'},{'kind':'code128','formula':'123456789012'}]}
    c=compose_labels(META,{'headers':HEADERS,'fields':[0],'records':RECORDS,'options':options})
    lines=c['pages'][0][0]['text']['richLines']
    assert any(line['kind']=='qr' and line['imageData'] for line in lines)
    assert any(line['kind']=='code128' and line['imageData'] for line in lines)
    assert any(s.get('bold') for line in lines for s in line['segments'])
    assert export_docx(TEMPLATE,META,c,options)

def test_group_page_break():
    records=[RECORDS[0],{'id':3,'qty':1,'values':['B','Atatürk','Konak','İzmir','05XX','2/2','100']}]
    options={'groupField':3,'groupNewPage':True}
    c=compose_labels(META,{'headers':HEADERS,'fields':[0],'records':records,'options':options})
    assert c['total']==2 and c['pageCount']==2 and c['groupBlankCount']>0

def test_check():
    s=validate_records(HEADERS,[RECORDS[0],{'id':3,'qty':0,'values':RECORDS[0]['values']}],[0],{})
    assert s['checked']==2 and s['summary']['duplicate']==1 and s['summary']['zero']==1
