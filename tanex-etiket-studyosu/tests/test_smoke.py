"""Stand-alone regression checks: python tests/test_smoke.py"""
import sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from docx import Document
from io import BytesIO
from core import inspect_template,compose_labels,export_docx

for tfile in ROOT.glob('templates/*.docx'):
    original=tfile.read_bytes()
    meta=inspect_template(original,tfile.name)
    assert meta['perPage']>0
    request={
      'fields':[0,1],'headers':['Kütüphane Adı','Adres'],
      'records':[{'id':i+2,'qty':1,'values':[f'Test Kütüphanesi {i+1}','Ankara / Türkiye']} for i in range(meta['perPage']+2)],
      'start':1,'options':{'font':'Arial','padding':3,'minFontSize':6,'maxFontSize':26}
    }
    built=compose_labels(meta,request)
    assert built['pageCount']==2
    out=export_docx(original,meta,built,request['options'])
    doc=Document(BytesIO(out))
    assert len(doc.tables)==2
    assert 'Test Kütüphanesi 1' in '\n'.join(c.text for row in doc.tables[0].rows for c in row.cells)
    print('OK:',tfile.name,'etiket:',meta['perPage'],'çıktı sayfa:',built['pageCount'])
print('Tüm temel testler geçti.')
