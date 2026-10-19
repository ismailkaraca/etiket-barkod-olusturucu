"""Tanex Etiket Stüdyosu — DOCX template inspection, Excel import and print layout.

Local-only application. No external service, database or telemetry.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from collections import Counter
import math
import re
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import ImageFont

MAX_LABELS = 600
FONT_DIR = Path('/usr/share/fonts/truetype')
FONT_FILES = {
    'Arial': ['/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'],
    'Calibri': ['/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'],
    'Aptos': ['/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'],
    'Times New Roman': ['/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf'],
}

def _font(pt: float, name='Arial'):
    p = next((f for f in FONT_FILES.get(name, FONT_FILES['Arial']) if Path(f).exists()),'/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    try:return ImageFont.truetype(p, max(6,round(pt*4)))  # Approximate measuring scale
    except OSError:return ImageFont.load_default(size=max(6,round(pt*4)))

def _unique_table_cells(tbl):
    seen = set(); out=[]
    for row_index, row in enumerate(tbl.rows):
        for col_index,cell in enumerate(row.cells):
            key = id(cell._tc)
            if key in seen: continue
            seen.add(key)
            out.append((row_index,col_index,cell))
    return out

def inspect_template(data: bytes, name: str) -> dict:
    if len(data)>8_000_000:raise ValueError('Word dosyası 8 MB sınırını aşıyor.')
    try: doc=Document(BytesIO(data))
    except Exception:raise ValueError('Geçerli bir .docx dosyası yükleyin.')
    if len(doc.sections)!=1:raise ValueError('İlk sürüm yalnızca tek bölümden oluşan Word şablonlarını destekliyor.')
    if len(doc.tables)!=1:raise ValueError('Word dosyası tam olarak bir ana etiket tablosu içermeli.')
    if doc.element.body.xpath('.//w:drawing') or doc.element.body.xpath('.//w:pict') or doc.element.body.xpath('.//w:txbxContent'):
        raise ValueError('Görsel, şekil veya metin kutusu içeren şablonlar bu sürümde desteklenmiyor.')
    table=doc.tables[0]; section=doc.sections[0]
    if not table.rows or not table.columns:raise ValueError('Şablonda etiket hücresi bulunamadı.')
    if any(c._tc.xpath('.//w:txbxContent') for _,_,c in _unique_table_cells(table)):
        raise ValueError('Metin kutusu içeren tablolar şu anda desteklenmiyor.')
    if any(c.tables for _,_,c in _unique_table_cells(table)):
        raise ValueError('İç içe tablolar şu anda desteklenmiyor.')
    pagew=round(section.page_width.mm,3); pageh=round(section.page_height.mm,3)
    m={'top':round((section.top_margin or Mm(0)).mm,3), 'bottom':round((section.bottom_margin or Mm(0)).mm,3),'left':round((section.left_margin or Mm(0)).mm,3),'right':round((section.right_margin or Mm(0)).mm,3)}
    # Most commercial label templates have explicit fixed row heights, and fixed grid column widths.
    widths = [round(table.cell(0,i).width.mm,2) for i in range(len(table.columns))]
    heights = [round((row.height or Mm((pageh-m['top']-m['bottom'])/len(table.rows))).mm,2) for row in table.rows]
    cells=[]
    for r,c,cell in _unique_table_cells(table):
        cells.append({'index':len(cells),'r':r,'c':c,'width':round(cell.width.mm,2),'height':heights[r], 'text':cell.text.strip()[:40]})
    if not cells:raise ValueError('Şablonda kullanılabilir hücre bulunamadı.')
    if len(cells)>600:raise ValueError('Bir sayfada en fazla 600 etiket hücresi destekleniyor.')
    return {'name':name,'page':{'width':pagew,'height':pageh},'margins':m,'rows':len(table.rows),'cols':len(table.columns),'widths':widths,'heights':heights,'cells':cells,'perPage':len(cells), 'warnings':(['Bu dosya A4 değil; Word belgesinin özgün sayfa boyutu kullanılacak.'] if not (abs(pagew-210)<3 and abs(pageh-297)<3) else [])}

def to_text(value, heading=''):
    if value is None:return ''
    if isinstance(value,bool):return 'Evet' if value else 'Hayır'
    if isinstance(value,(int,float)):
        if heading and ('tarih' in heading.casefold() or 'zaman' in heading.casefold()) and 20000<float(value)<100000:
            dt=datetime(1899,12,30)+timedelta(days=float(value))
            return dt.strftime('%d.%m.%Y %H:%M') if dt.hour or dt.minute else dt.strftime('%d.%m.%Y')
        if isinstance(value,float) and value.is_integer():return str(int(value))
    return str(value).strip()

def inspect_xlsx(file_path: str) -> dict:
    """Read workbook values using the .xlsx Open XML standard (Python stdlib).

    Works offline without a proprietary spreadsheet runtime. Cached formula results
    are used; formulas without cached results appear blank.
    """
    import zipfile
    import posixpath
    from xml.etree import ElementTree as ET

    XML_NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    REL_NS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    PKG_NS='http://schemas.openxmlformats.org/package/2006/relationships'
    M=lambda tag:'{%s}%s' % (XML_NS,tag)
    P=lambda tag:'{%s}%s' % (PKG_NS,tag)
    try:z=zipfile.ZipFile(file_path)
    except Exception as exc:raise ValueError('Dosya XLSX/ZIP biçiminde okunamadı.') from exc
    with z:
        namelist=set(z.namelist())
        if len(namelist)>10000:raise ValueError('Excel dosyası çok fazla dahili dosya içeriyor.')
        def loadxml(name):
            if name not in namelist:raise ValueError('Eksik XLSX bileşeni: '+name)
            if z.getinfo(name).file_size>30_000_000:raise ValueError('XLSX bileşeni çok büyük: '+name)
            return ET.fromstring(z.read(name))
        book=loadxml('xl/workbook.xml')
        relroot=loadxml('xl/_rels/workbook.xml.rels')
        rels={x.attrib.get('Id'):x.attrib.get('Target','') for x in relroot.findall(P('Relationship'))}
        shared=[]
        if 'xl/sharedStrings.xml' in namelist:
            root=loadxml('xl/sharedStrings.xml')
            for si in root.findall(M('si')):
                shared.append(''.join(el.text or '' for el in si.iter(M('t'))))
        styles=[];fmt_map={}
        if 'xl/styles.xml' in namelist:
            style_root=loadxml('xl/styles.xml')
            numfmts=style_root.find(M('numFmts'))
            if numfmts is not None:
                fmt_map={int(x.attrib['numFmtId']):x.attrib.get('formatCode','') for x in numfmts.findall(M('numFmt'))}
            xfs=style_root.find(M('cellXfs'))
            if xfs is not None:styles=[int(x.attrib.get('numFmtId',0)) for x in xfs.findall(M('xf'))]
        sheets=[]
        for spec in book.findall('.//'+M('sheets')+'/'+M('sheet')):
            name=spec.attrib.get('name','Sayfa')
            target=rels.get(spec.attrib.get('{%s}id'%REL_NS,''),'')
            if not target:continue
            if target.startswith('/'):path=target.lstrip('/')
            else:path=posixpath.normpath(posixpath.join('xl',target))
            if path not in namelist:continue
            tree=loadxml(path)
            sheetdata=tree.find(M('sheetData'))
            if sheetdata is None:continue
            rawrows=[]
            for item in sheetdata.findall(M('row')):
                rowid=int(item.attrib.get('r',len(rawrows)+1))
                if rowid>6001:break
                cells={}
                for c in item.findall(M('c')):
                    ref=c.attrib.get('r','')
                    chars=re.match(r'([A-Z]+)',ref)
                    if not chars:continue
                    col=0
                    for ch in chars.group(1):col=col*26+ord(ch)-64
                    col-=1
                    if col>=120:continue
                    typ=c.attrib.get('t','n')
                    value=c.find(M('v'))
                    if typ=='inlineStr':
                        istr=c.find(M('is'))
                        val=''.join(t.text or '' for t in istr.iter(M('t'))) if istr is not None else ''
                    elif value is None:val=None
                    elif typ=='s':
                        try:val=shared[int(value.text)]
                        except (ValueError,IndexError,TypeError):val=''
                    elif typ=='b':val=value.text=='1'
                    elif typ in ('str','e'):val=value.text or ''
                    else:
                        try:
                            v=float(value.text)
                            val=int(v) if v.is_integer() else v
                            fmtid=styles[int(c.attrib.get('s',0))] if styles else 0
                            custom=fmt_map.get(fmtid,'')
                            if re.fullmatch(r'0{2,}',custom) and isinstance(val,int):val=str(val).zfill(len(custom))
                            if fmtid in {14,15,16,17,18,19,20,21,22,45,46,47} or (custom and ('dd' in custom.lower() or 'yy' in custom.lower())):
                                if 1000 < float(val) < 100000:
                                    dt=datetime(1899,12,30)+timedelta(days=float(val))
                                    val=dt.strftime('%d.%m.%Y %H:%M') if dt.hour or dt.minute else dt.strftime('%d.%m.%Y')
                        except (ValueError,TypeError,OverflowError):val=value.text
                    cells[col]=val
                rawrows.append((rowid,cells))
            if not rawrows:continue
            headers0=rawrows[0][1]
            maxcol=min(max((c for c,v in headers0.items() if v is not None and str(v).strip()),default=-1)+1,120)
            if maxcol==0:continue
            dup=Counter();heads=[]
            for c in range(maxcol):
                v=headers0.get(c)
                label=str(v).strip() if v is not None else f'Sütun {c+1}'
                if not label:label=f'Sütun {c+1}'
                dup[label]+=1
                if dup[label]>1:label=f'{label} ({dup[label]})'
                heads.append(label)
            data=[]
            for rowid,cells in rawrows[1:]:
                if not any(v is not None and str(v).strip() for v in cells.values()):continue
                data.append({'id':rowid,'values':[to_text(cells.get(c),heads[c]) for c in range(maxcol)]})
            sheets.append({'name':name,'columns':heads,'rows':data,'count':len(data)})
    if not sheets:raise ValueError('Excel dosyasında okunabilir veri bulunamadı.')
    return {'sheets':sheets}

def _text_width_mm(s, font):
    return font.getlength(s) / 4.0 * 25.4/72

def wrap_text(text: str, max_mm:float, font):
    """Preserve manual line breaks; character split only for exceptionally long words."""
    result=[]
    for paragraph in (text.split('\n') if text else ['']):
        words=paragraph.split(' ')
        line=''
        for word in words:
            trial=(line+' '+word).strip() if line else word
            if not line or _text_width_mm(trial,font)<=max_mm:
                line=trial
            else:
                result.append(line)
                line=word
            while line and _text_width_mm(line,font)>max_mm:
                cut=len(line)
                while cut>1 and _text_width_mm(line[:cut],font)>max_mm:cut-=1
                result.append(line[:cut]);line=line[cut:]
        result.append(line)
    return result

def _layout_text(parts, width_mm,height_mm, options):
    padding=max(0,float(options.get('padding',3)))
    width=max(2,width_mm-2*padding)
    height=max(2,height_mm-2*padding)
    name=options.get('font','Arial')
    label_names=options.get('showNames',False)
    base=float(options.get('fontSize',16))
    minpt=max(4,float(options.get('minFontSize',6)))
    maxpt=max(minpt,float(options.get('maxFontSize',26)))
    manual=bool(options.get('manual',False))
    strings=[((heading+': ') if label_names else '')+str(val) for heading,val in parts if str(val).strip()]
    if not strings:strings=['']
    if manual: sizes=[max(4,min(base,72))]
    else: sizes=[k/2 for k in range(round(maxpt*2),round(minpt*2)-1,-1)]
    candidate=None
    for size in sizes:
        f=_font(size,name)
        lines=[]
        for j,txt in enumerate(strings):
            if j>0:lines.append('') if options.get('fieldBlankLine',False) else None
            lines.extend(wrap_text(txt,width,f))
        # Word's actual metrics depend on font, line-height and paragraph defaults.
        h=len(lines)*size*25.4/72*1.23
        over=h>height+0.05 or any(_text_width_mm(l,f)>width+0.1 for l in lines)
        candidate={'size':size,'lines':lines,'overflow':over,'heightMM':round(h,1),'usableHeightMM':round(height,1),'usableWidthMM':round(width,1)}
        if not over:break
    return candidate

def compose_labels(template, request):
    fields=[int(i) for i in request.get('fields',[])]
    if not fields:raise ValueError('Yazdırılacak en az bir sütun seçmelisin.')
    heads=request.get('headers',[]); records=request.get('records',[])
    if len(fields)>len(heads) or any(i<0 or i>=len(heads) for i in fields):raise ValueError('Sütun bilgileri hatalı.')
    first=int(request.get('start',0))
    if first<0 or first>=template['perPage']:raise ValueError('Başlangıç etiketi geçersiz.')
    opts=request.get('options',{})
    out=[]
    for rec in records:
        copies=int(rec.get('qty',1))
        if not 0<=copies<=200:raise ValueError('Bir satır en fazla 200 defa çoğaltılabilir.')
        values=rec.get('values',[])
        parts=[(heads[c],values[c] if c<len(values) else '') for c in fields]
        for _ in range(copies):out.append({'sourceRow':rec.get('id'),'parts':parts})
        if len(out)>MAX_LABELS:raise ValueError(f'En fazla {MAX_LABELS} etiket tek seferde oluşturulabilir.')
    total=first+len(out)
    if not out:raise ValueError('En az bir satır ve etiket adedi seçmelisin.')
    per=template['perPage']; pages=[]; overflows=0
    for pg in range(math.ceil(total/per)):
        labels=[]
        for k,geom in enumerate(template['cells']):
            idx=pg*per+k-first
            if not 0<=idx<len(out):labels.append(None);continue
            item=out[idx]
            info=_layout_text(item['parts'],geom['width'],geom['height'],opts)
            if info['overflow']:overflows+=1
            labels.append({'sourceRow':item['sourceRow'],'position':k+1,'text':info,'parts':item['parts']})
        pages.append(labels)
    return {'pages':pages,'total':len(out),'pageCount':len(pages),'overflowCount':overflows,'start':first}

def _replace_contents(cell, lines, font_size, options):
    # Keep original table cell properties (borders/widths/vertical merges). Replace content only.
    tc=cell._tc
    for node in list(tc):
        if node.tag != qn('w:tcPr'):tc.remove(node)
    padding=max(0,float(options.get('padding',3)))
    props=tc.get_or_add_tcPr()
    mar=props.find(qn('w:tcMar'))
    if mar is None:mar=OxmlElement('w:tcMar');props.append(mar)
    for side in ['top','left','bottom','right']:
        e=mar.find(qn('w:'+side))
        if e is None:e=OxmlElement('w:'+side);mar.append(e)
        e.set(qn('w:w'),str(round(padding*56.692913))) # mm to 1/20pt
        e.set(qn('w:type'),'dxa')
    valign=options.get('vertical','center')
    cell.vertical_alignment={'top':WD_CELL_VERTICAL_ALIGNMENT.TOP,'center':WD_CELL_VERTICAL_ALIGNMENT.CENTER,'bottom':WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}.get(valign,WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    halign={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}.get(options.get('align','center'),WD_ALIGN_PARAGRAPH.CENTER)
    p=OxmlElement('w:p'); tc.append(p)
    from docx.text.paragraph import Paragraph
    para=Paragraph(p,cell)
    para.alignment=halign
    pf=para.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.0
    for i,line in enumerate(lines):
        if i:para.add_run().add_break()
        run=para.add_run(line)
        run.font.name=options.get('font','Arial')
        run.font.size=Pt(font_size)
        run.font.bold=bool(options.get('bold',False))

def export_docx(template_data:bytes,template,composition,options):
    doc=Document(BytesIO(template_data))
    table=doc.tables[0]
    # Remove empty spacer paragraphs supplied by blank commercial label templates.
    # On a physically full page, these would cause phantom extra pages.
    for paragraph in list(doc.paragraphs):
        if not paragraph.text and not paragraph._p.xpath('.//w:drawing') and not paragraph._p.xpath('.//w:pict'):
            paragraph._p.getparent().remove(paragraph._p)
    page_count=len(composition['pages'])
    source_xml=deepcopy(table._tbl)
    # The first template page stays in place; subsequent page copies retain its precise geometry.
    for pnum,page in enumerate(composition['pages']):
        if pnum:
            p=doc.add_paragraph()
            p.paragraph_format.page_break_before=True
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=Pt(0.1)
            r=p.add_run()
            r.font.size=Pt(1)
            tb=deepcopy(source_xml)
            p._p.addnext(tb)
            from docx.table import Table
            tbl=Table(tb,doc._body)
        else:tbl=table
        cells=_unique_table_cells(tbl)
        for index,(_r,_c,cell) in enumerate(cells):
            obj=page[index] if index<len(page) else None
            if obj:
                _replace_contents(cell,obj['text']['lines'],obj['text']['size'],options)
            else:
                _replace_contents(cell,[],float(options.get('fontSize',16)),options)
    # Prevent blank paragraph at end of the sheet from becoming a phantom extra page.
    for para in doc.paragraphs:
        if not para.text and not para._p.xpath('.//w:br') and not para._p.xpath('.//w:pPr/w:pageBreakBefore'):
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0)
            para.paragraph_format.line_spacing=1
            for r in para.runs:r.font.size=Pt(1)
    output=BytesIO();doc.save(output);return output.getvalue()
